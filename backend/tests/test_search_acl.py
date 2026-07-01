from __future__ import annotations

from collections.abc import Callable
from io import BytesIO
from uuid import UUID, uuid4

import pytest
from docx import Document
from httpx import AsyncClient
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker

from app.core.config import Settings
from app.infrastructure.search.testing import InMemorySearchIndexAdapter
from app.infrastructure.storage.testing import InMemoryStorageAdapter
from app.modules.audit.dispatcher import (
    LoggingOutboxPublisher,
    OutboxDispatcher,
    OutboxDispatchResult,
)
from app.modules.audit.models import OutboxEvent
from app.modules.audit.repository import AuditRepository
from app.modules.audit.service import AuditService
from app.modules.file.models import FileBlob, FileVersion
from app.modules.permission.events import emit_permission_changed
from app.modules.permission.models import AclEntry, SpaceMember
from app.modules.search.acl import build_index_acl_token_set, build_index_acl_tokens
from app.modules.search.events import (
    SEARCH_ACL_REBUILD_REQUESTED,
    SEARCH_EXTRACT_REQUESTED,
    SEARCH_INDEX_REQUESTED,
)
from app.modules.search.extractor import SearchExtractionService
from app.modules.search.extractors import (
    DocxTextExtractor,
    PdfTextExtractor,
    PptxTextExtractor,
    TextExtractionContext,
    TextExtractionError,
    TextExtractor,
    XlsxTextExtractor,
)
from app.modules.search.indexer import SearchIndexService
from app.modules.search.repository import SearchRepository
from app.workers import search_tasks
from tests.helpers import client as client
from tests.helpers import create_space, login, seed_admin
from tests.helpers import session_factory as session_factory
from tests.helpers import settings as settings
from tests.helpers import storage_adapter as storage_adapter
from tests.test_permission_cache import _get_seeded_tenant_and_user


class FailingReadStorageAdapter(InMemoryStorageAdapter):
    async def read_object_bytes(
        self,
        *,
        bucket: str,
        storage_key: str,
        max_bytes: int,
    ) -> bytes:
        raise RuntimeError("storage unavailable")


class OversizedTextExtractor:
    def supports(self, context: TextExtractionContext) -> bool:
        return True

    def extract(self, content: bytes, context: TextExtractionContext) -> str:
        return "x" * 9


def test_pdf_text_extractor_uses_pypdf_for_text_content() -> None:
    pdf_bytes = _build_pdf_with_text("Hello PDF Search")

    text = PdfTextExtractor().extract(
        pdf_bytes,
        TextExtractionContext(mime_type="application/pdf", name="document.pdf"),
    )

    assert text == "Hello PDF Search"


def test_pdf_text_extractor_limits_pages() -> None:
    pdf_bytes = _build_pdf_with_text("First Page", "Second Page")

    text = PdfTextExtractor(max_pages=1).extract(
        pdf_bytes,
        TextExtractionContext(mime_type="application/pdf", name="document.pdf"),
    )

    assert text == "First Page"


def test_docx_text_extractor_reads_paragraphs_and_table_cells() -> None:
    docx_bytes = _build_docx_with_text("Docx Paragraph", ("Cell A", "Cell B"))

    text = DocxTextExtractor().extract(
        docx_bytes,
        TextExtractionContext(
            mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            name="document.docx",
        ),
    )

    assert text == "Docx Paragraph\nCell A\nCell B"


def test_pptx_text_extractor_reads_shapes_and_table_cells() -> None:
    pptx_bytes = _build_pptx_with_text("PPTX Shape", ("Slide Cell A", "Slide Cell B"))

    text = PptxTextExtractor().extract(
        pptx_bytes,
        TextExtractionContext(
            mime_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
            name="slides.pptx",
        ),
    )

    assert text == "PPTX Shape\nSlide Cell A\nSlide Cell B"


def test_xlsx_text_extractor_reads_cell_values() -> None:
    xlsx_bytes = _build_xlsx_with_text("XLSX Cell", ("Sheet Cell A", "Sheet Cell B"))

    text = XlsxTextExtractor().extract(
        xlsx_bytes,
        TextExtractionContext(
            mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            name="sheet.xlsx",
        ),
    )

    assert text == "XLSX Cell\nSheet Cell A\nSheet Cell B"


@pytest.mark.parametrize(
    ("extractor", "content", "context", "reason"),
    [
        (
            PptxTextExtractor(max_entries=1),
            lambda: _build_pptx_with_text("Large PPTX", ("Cell A", "Cell B")),
            TextExtractionContext(
                mime_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                name="slides.pptx",
            ),
            "pptx_archive_too_large",
        ),
        (
            XlsxTextExtractor(max_entries=1),
            lambda: _build_xlsx_with_text("Large XLSX", ("Cell A", "Cell B")),
            TextExtractionContext(
                mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                name="sheet.xlsx",
            ),
            "xlsx_archive_too_large",
        ),
    ],
)
def test_ooxml_text_extractors_skip_when_archive_exceeds_limit(
    extractor: TextExtractor,
    content: Callable[[], bytes],
    context: TextExtractionContext,
    reason: str,
) -> None:
    with pytest.raises(TextExtractionError) as exc_info:
        extractor.extract(content(), context)

    assert exc_info.value.status == "skipped"
    assert exc_info.value.reason == reason


def test_build_index_acl_tokens_includes_roles_and_allowed_subjects() -> None:
    tenant_id = uuid4()
    space_id = uuid4()
    owner_id = uuid4()
    viewer_id = uuid4()
    node_id = uuid4()
    department_id = uuid4()
    group_id = uuid4()
    tokens = build_index_acl_tokens(
        space_members=[
            SpaceMember(
                tenant_id=tenant_id,
                space_id=space_id,
                user_id=owner_id,
                role="owner",
                created_by=None,
            ),
            SpaceMember(
                tenant_id=tenant_id,
                space_id=space_id,
                user_id=viewer_id,
                role="viewer",
                created_by=owner_id,
            ),
        ],
        acl_entries=[
            AclEntry(
                tenant_id=tenant_id,
                node_id=node_id,
                subject_type="department",
                subject_id=department_id,
                effect="allow",
                actions=["download"],
                inherit=True,
                created_by=owner_id,
            ),
            AclEntry(
                tenant_id=tenant_id,
                node_id=node_id,
                subject_type="group",
                subject_id=group_id,
                effect="deny",
                actions=["download"],
                inherit=True,
                created_by=owner_id,
            ),
        ],
    )

    assert tokens == sorted(
        [
            f"space:{space_id}:role:owner",
            f"space:{space_id}:role:viewer",
            f"department:{department_id}",
        ]
    )


def test_build_index_acl_tokens_ignores_non_visible_and_denied_acl() -> None:
    tenant_id = uuid4()
    node_id = uuid4()
    owner_id = uuid4()
    upload_only_user_id = uuid4()
    denied_user_id = uuid4()

    tokens = build_index_acl_tokens(
        space_members=[],
        acl_entries=[
            AclEntry(
                tenant_id=tenant_id,
                node_id=node_id,
                subject_type="user",
                subject_id=upload_only_user_id,
                effect="allow",
                actions=["upload"],
                inherit=True,
                created_by=owner_id,
            ),
            AclEntry(
                tenant_id=tenant_id,
                node_id=node_id,
                subject_type="user",
                subject_id=denied_user_id,
                effect="allow",
                actions=["download"],
                inherit=True,
                created_by=owner_id,
            ),
            AclEntry(
                tenant_id=tenant_id,
                node_id=node_id,
                subject_type="user",
                subject_id=denied_user_id,
                effect="deny",
                actions=["download"],
                inherit=True,
                created_by=owner_id,
            ),
        ],
    )

    assert tokens == []


def test_build_index_acl_token_set_returns_deny_tokens_for_query_exclusion() -> None:
    tenant_id = uuid4()
    space_id = uuid4()
    node_id = uuid4()
    owner_id = uuid4()
    denied_group_id = uuid4()

    token_set = build_index_acl_token_set(
        space_members=[
            SpaceMember(
                tenant_id=tenant_id,
                space_id=space_id,
                user_id=owner_id,
                role="owner",
                created_by=None,
            )
        ],
        acl_entries=[
            AclEntry(
                tenant_id=tenant_id,
                node_id=node_id,
                subject_type="group",
                subject_id=denied_group_id,
                effect="deny",
                actions=["download"],
                inherit=True,
                created_by=owner_id,
            )
        ],
    )

    assert token_set.allow_tokens == [f"space:{space_id}:role:owner"]
    assert token_set.deny_tokens == [f"group:{denied_group_id}"]


@pytest.mark.asyncio
async def test_search_dispatcher_consumes_acl_rebuild_events(
    session_factory: async_sessionmaker[AsyncSession],
    settings: Settings,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    await seed_admin(session_factory, settings)
    tenant_id, _ = await _get_seeded_tenant_and_user(session_factory)
    node_id = uuid4()
    async with session_factory() as session:
        repository = AuditRepository(session)
        search_event = await repository.add_outbox_event(
            tenant_id=tenant_id,
            event_type=SEARCH_ACL_REBUILD_REQUESTED,
            aggregate_type="node",
            aggregate_id=node_id,
            payload={
                "scope": "node",
                "resource_id": str(node_id),
                "permission_version": 2,
                "reason": "node_acl_created",
            },
        )
        permission_event = await repository.add_outbox_event(
            tenant_id=tenant_id,
            event_type="permission.changed",
            aggregate_type="node",
            aggregate_id=node_id,
            payload={
                "scope": "node",
                "resource_id": str(node_id),
                "permission_version": 2,
            },
        )
        await session.commit()

    monkeypatch.setattr(search_tasks, "get_settings", lambda: settings)
    monkeypatch.setattr(search_tasks, "get_session_factory", lambda: session_factory)

    result = await search_tasks._dispatch_search_outbox(batch_size=10)

    assert result == {"claimed": 1, "sent": 1, "failed": 0, "dead": 0}
    async with session_factory() as session:
        stored_search_event = await session.get(OutboxEvent, search_event.id)
        stored_permission_event = await session.get(OutboxEvent, permission_event.id)

    assert stored_search_event is not None
    assert stored_search_event.status == "sent"
    assert stored_permission_event is not None
    assert stored_permission_event.status == "pending"


@pytest.mark.asyncio
async def test_permission_change_writes_search_acl_rebuild_event(
    session_factory: async_sessionmaker[AsyncSession],
    settings: Settings,
) -> None:
    await seed_admin(session_factory, settings)
    tenant_id, user_id = await _get_seeded_tenant_and_user(session_factory)
    node_id = uuid4()
    async with session_factory() as session:
        await emit_permission_changed(
            audit_service=AuditService(repository=AuditRepository(session)),
            tenant_id=tenant_id,
            actor_id=user_id,
            scope="node",
            resource_id=node_id,
            permission_version=3,
            reason="node_acl_created",
            affected_user_id=user_id,
            metadata={"space_id": str(uuid4()), "subject_type": "user"},
        )
        await session.commit()

    async with session_factory() as session:
        events = (
            await session.execute(
                select(OutboxEvent.event_type, OutboxEvent.payload).order_by(OutboxEvent.created_at)
            )
        ).all()

    assert [event_type for event_type, _ in events] == [
        "permission.changed",
        SEARCH_ACL_REBUILD_REQUESTED,
    ]
    assert events[1][1]["reason"] == "node_acl_created"


@pytest.mark.asyncio
async def test_instant_upload_writes_search_index_event(
    client: AsyncClient,
    session_factory: async_sessionmaker[AsyncSession],
    settings: Settings,
) -> None:
    await seed_admin(session_factory, settings)
    token = await login(client)
    space = await create_space(client, token, slug="search-index-event-space")
    tenant_id = UUID(str(space["tenant_id"]))
    async with session_factory() as session:
        session.add(
            FileBlob(
                tenant_id=tenant_id,
                hash_algo="sha256",
                content_hash="c" * 64,
                size_bytes=128,
                storage_key="objects/test/cc",
                mime_type="text/plain",
                ref_count=0,
            )
        )
        await session.commit()

    response = await client.post(
        "/api/v1/uploads/init",
        headers={"X-CSRF-Token": token},
        json={
            "space_id": space["id"],
            "parent_id": space["root_node_id"],
            "file_name": "搜索事件.txt",
            "size_bytes": 128,
            "content_hash": "c" * 64,
            "hash_algo": "sha256",
            "mime_type": "text/plain",
        },
    )

    assert response.status_code == 201
    node_id = response.json()["node_id"]
    async with session_factory() as session:
        event = (
            await session.execute(
                select(OutboxEvent).where(OutboxEvent.event_type == SEARCH_INDEX_REQUESTED)
            )
        ).scalar_one()

    assert event.aggregate_type == "node"
    assert event.aggregate_id == UUID(node_id)
    assert event.payload["node_id"] == node_id
    assert event.payload["reason"] == "upload_instant"

    async with session_factory() as session:
        extract_event = (
            await session.execute(
                select(OutboxEvent).where(OutboxEvent.event_type == SEARCH_EXTRACT_REQUESTED)
            )
        ).scalar_one()

    assert extract_event.aggregate_type == "file_version"
    assert extract_event.payload["node_id"] == node_id
    assert extract_event.payload["reason"] == "upload_instant"


@pytest.mark.asyncio
async def test_search_index_requested_event_indexes_file_document(
    client: AsyncClient,
    session_factory: async_sessionmaker[AsyncSession],
    settings: Settings,
) -> None:
    await seed_admin(session_factory, settings)
    token = await login(client)
    space = await create_space(client, token, slug="search-index-worker-space")
    tenant_id = UUID(str(space["tenant_id"]))
    async with session_factory() as session:
        session.add(
            FileBlob(
                tenant_id=tenant_id,
                hash_algo="sha256",
                content_hash="d" * 64,
                size_bytes=256,
                storage_key="objects/test/dd",
                mime_type="text/plain",
                ref_count=0,
            )
        )
        await session.commit()

    response = await client.post(
        "/api/v1/uploads/init",
        headers={"X-CSRF-Token": token},
        json={
            "space_id": space["id"],
            "parent_id": space["root_node_id"],
            "file_name": "搜索索引.txt",
            "size_bytes": 256,
            "content_hash": "d" * 64,
            "hash_algo": "sha256",
            "mime_type": "text/plain",
        },
    )
    assert response.status_code == 201
    node_id = response.json()["node_id"]
    index_adapter = InMemorySearchIndexAdapter()

    async with session_factory() as session:
        dispatcher = OutboxDispatcher(
            repository=AuditRepository(session),
            publisher=search_tasks.SearchOutboxPublisher(
                index_service=SearchIndexService(
                    repository=SearchRepository(session),
                    index_adapter=index_adapter,
                ),
                fallback_publisher=LoggingOutboxPublisher(),
            ),
            max_retries=settings.outbox_max_retries,
        )
        result = await dispatcher.dispatch_pending(
            batch_size=10,
            event_types=[SEARCH_INDEX_REQUESTED],
        )
        await session.commit()

    assert result.to_dict() == {"claimed": 1, "sent": 1, "failed": 0, "dead": 0}
    document = index_adapter.documents[f"{tenant_id}:{node_id}"]
    assert document.name == "搜索索引.txt"
    assert document.node_id == node_id
    assert document.space_id == str(space["id"])
    assert document.content_hash == "d" * 64
    assert document.acl_tokens == [f"space:{space['id']}:role:owner"]
    assert document.deny_acl_tokens == []


@pytest.mark.asyncio
async def test_search_extract_requested_indexes_text_content(
    client: AsyncClient,
    session_factory: async_sessionmaker[AsyncSession],
    settings: Settings,
    storage_adapter: InMemoryStorageAdapter,
) -> None:
    await seed_admin(session_factory, settings)
    token = await login(client)
    space = await create_space(client, token, slug="search-text-extract-space")
    tenant_id = UUID(str(space["tenant_id"]))
    storage_key = "objects/test/text-content"
    storage_adapter.object_contents[(settings.s3_bucket, storage_key)] = (
        "标题\n可搜索正文内容".encode()
    )
    async with session_factory() as session:
        session.add(
            FileBlob(
                tenant_id=tenant_id,
                hash_algo="sha256",
                content_hash="1" * 64,
                size_bytes=128,
                storage_key=storage_key,
                mime_type="text/plain",
                ref_count=0,
            )
        )
        await session.commit()

    response = await client.post(
        "/api/v1/uploads/init",
        headers={"X-CSRF-Token": token},
        json={
            "space_id": space["id"],
            "parent_id": space["root_node_id"],
            "file_name": "正文.txt",
            "size_bytes": 128,
            "content_hash": "1" * 64,
            "hash_algo": "sha256",
            "mime_type": "text/plain",
        },
    )
    assert response.status_code == 201
    node_id = response.json()["node_id"]
    version_id = UUID(response.json()["version_id"])
    index_adapter = InMemorySearchIndexAdapter()

    await _dispatch_search_events(
        session_factory=session_factory,
        settings=settings,
        index_adapter=index_adapter,
        storage_adapter=storage_adapter,
        batch_size=10,
        event_types=[SEARCH_EXTRACT_REQUESTED],
    )

    async with session_factory() as session:
        version = await session.get(FileVersion, version_id)
    assert version is not None
    assert version.search_status == "indexed"
    assert version.search_text == "标题\n可搜索正文内容"
    assert version.search_error is None
    assert index_adapter.documents[f"{tenant_id}:{node_id}"].content == "标题\n可搜索正文内容"


@pytest.mark.asyncio
async def test_search_extract_requested_indexes_pdf_text_content(
    client: AsyncClient,
    session_factory: async_sessionmaker[AsyncSession],
    settings: Settings,
    storage_adapter: InMemoryStorageAdapter,
) -> None:
    await seed_admin(session_factory, settings)
    token = await login(client)
    space = await create_space(client, token, slug="search-pdf-extract-space")
    tenant_id = UUID(str(space["tenant_id"]))
    storage_key = "objects/test/pdf-content"
    pdf_bytes = _build_pdf_with_text("PDF Searchable Body")
    storage_adapter.object_contents[(settings.s3_bucket, storage_key)] = pdf_bytes
    async with session_factory() as session:
        session.add(
            FileBlob(
                tenant_id=tenant_id,
                hash_algo="sha256",
                content_hash="6" * 64,
                size_bytes=len(pdf_bytes),
                storage_key=storage_key,
                mime_type="application/pdf",
                ref_count=0,
            )
        )
        await session.commit()

    response = await client.post(
        "/api/v1/uploads/init",
        headers={"X-CSRF-Token": token},
        json={
            "space_id": space["id"],
            "parent_id": space["root_node_id"],
            "file_name": "可搜索PDF.pdf",
            "size_bytes": len(pdf_bytes),
            "content_hash": "6" * 64,
            "hash_algo": "sha256",
            "mime_type": "application/pdf",
        },
    )
    assert response.status_code == 201
    node_id = response.json()["node_id"]
    version_id = UUID(response.json()["version_id"])
    index_adapter = InMemorySearchIndexAdapter()

    await _dispatch_search_events(
        session_factory=session_factory,
        settings=settings,
        index_adapter=index_adapter,
        storage_adapter=storage_adapter,
        batch_size=10,
        event_types=[SEARCH_EXTRACT_REQUESTED],
    )

    async with session_factory() as session:
        version = await session.get(FileVersion, version_id)
    assert version is not None
    assert version.search_status == "indexed"
    assert version.search_text == "PDF Searchable Body"
    assert version.search_error is None
    assert index_adapter.documents[f"{tenant_id}:{node_id}"].content == "PDF Searchable Body"


@pytest.mark.asyncio
async def test_search_extract_requested_indexes_docx_text_content(
    client: AsyncClient,
    session_factory: async_sessionmaker[AsyncSession],
    settings: Settings,
    storage_adapter: InMemoryStorageAdapter,
) -> None:
    await seed_admin(session_factory, settings)
    token = await login(client)
    space = await create_space(client, token, slug="search-docx-extract-space")
    tenant_id = UUID(str(space["tenant_id"]))
    storage_key = "objects/test/docx-content"
    docx_bytes = _build_docx_with_text("DOCX Searchable Body", ("表格一", "表格二"))
    storage_adapter.object_contents[(settings.s3_bucket, storage_key)] = docx_bytes
    async with session_factory() as session:
        session.add(
            FileBlob(
                tenant_id=tenant_id,
                hash_algo="sha256",
                content_hash="8" * 64,
                size_bytes=len(docx_bytes),
                storage_key=storage_key,
                mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                ref_count=0,
            )
        )
        await session.commit()

    response = await client.post(
        "/api/v1/uploads/init",
        headers={"X-CSRF-Token": token},
        json={
            "space_id": space["id"],
            "parent_id": space["root_node_id"],
            "file_name": "可搜索DOCX.docx",
            "size_bytes": len(docx_bytes),
            "content_hash": "8" * 64,
            "hash_algo": "sha256",
            "mime_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        },
    )
    assert response.status_code == 201
    node_id = response.json()["node_id"]
    version_id = UUID(response.json()["version_id"])
    index_adapter = InMemorySearchIndexAdapter()

    await _dispatch_search_events(
        session_factory=session_factory,
        settings=settings,
        index_adapter=index_adapter,
        storage_adapter=storage_adapter,
        batch_size=10,
        event_types=[SEARCH_EXTRACT_REQUESTED],
    )

    async with session_factory() as session:
        version = await session.get(FileVersion, version_id)
    assert version is not None
    assert version.search_status == "indexed"
    assert version.search_text == "DOCX Searchable Body\n表格一\n表格二"
    assert version.search_error is None
    assert index_adapter.documents[f"{tenant_id}:{node_id}"].content == (
        "DOCX Searchable Body\n表格一\n表格二"
    )


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("slug", "file_name", "mime_type", "content_hash", "content", "expected_text"),
    [
        (
            "search-pptx-extract-space",
            "可搜索PPTX.pptx",
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            "a" * 64,
            lambda: _build_pptx_with_text("PPTX Searchable Body", ("幻灯片表格一", "幻灯片表格二")),
            "PPTX Searchable Body\n幻灯片表格一\n幻灯片表格二",
        ),
        (
            "search-xlsx-extract-space",
            "可搜索XLSX.xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "b" * 64,
            lambda: _build_xlsx_with_text("XLSX Searchable Body", ("表格单元格一", "表格单元格二")),
            "XLSX Searchable Body\n表格单元格一\n表格单元格二",
        ),
    ],
)
async def test_search_extract_requested_indexes_ooxml_text_content(
    client: AsyncClient,
    session_factory: async_sessionmaker[AsyncSession],
    settings: Settings,
    storage_adapter: InMemoryStorageAdapter,
    slug: str,
    file_name: str,
    mime_type: str,
    content_hash: str,
    content: Callable[[], bytes],
    expected_text: str,
) -> None:
    tenant_id, node_id, version_id, index_adapter = await _upload_existing_blob_and_extract_text(
        client=client,
        session_factory=session_factory,
        settings=settings,
        storage_adapter=storage_adapter,
        slug=slug,
        storage_key=f"objects/test/{slug}",
        file_name=file_name,
        content_hash=content_hash,
        mime_type=mime_type,
        content=content(),
    )

    async with session_factory() as session:
        version = await session.get(FileVersion, version_id)
    assert version is not None
    assert version.search_status == "indexed"
    assert version.search_text == expected_text
    assert version.search_error is None
    assert index_adapter.documents[f"{tenant_id}:{node_id}"].content == expected_text


@pytest.mark.asyncio
async def test_search_extract_requested_skips_docx_when_archive_exceeds_limit(
    client: AsyncClient,
    session_factory: async_sessionmaker[AsyncSession],
    settings: Settings,
    storage_adapter: InMemoryStorageAdapter,
) -> None:
    await seed_admin(session_factory, settings)
    token = await login(client)
    space = await create_space(client, token, slug="search-docx-archive-limit-space")
    tenant_id = UUID(str(space["tenant_id"]))
    storage_key = "objects/test/docx-archive-limit"
    docx_bytes = _build_docx_with_text("Large Archive", ("Cell A", "Cell B"))
    storage_adapter.object_contents[(settings.s3_bucket, storage_key)] = docx_bytes
    async with session_factory() as session:
        session.add(
            FileBlob(
                tenant_id=tenant_id,
                hash_algo="sha256",
                content_hash="9" * 64,
                size_bytes=len(docx_bytes),
                storage_key=storage_key,
                mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                ref_count=0,
            )
        )
        await session.commit()

    response = await client.post(
        "/api/v1/uploads/init",
        headers={"X-CSRF-Token": token},
        json={
            "space_id": space["id"],
            "parent_id": space["root_node_id"],
            "file_name": "过大DOCX.docx",
            "size_bytes": len(docx_bytes),
            "content_hash": "9" * 64,
            "hash_algo": "sha256",
            "mime_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        },
    )
    assert response.status_code == 201
    version_id = UUID(response.json()["version_id"])
    index_adapter = InMemorySearchIndexAdapter()

    await _dispatch_search_events(
        session_factory=session_factory,
        settings=settings,
        index_adapter=index_adapter,
        storage_adapter=storage_adapter,
        batch_size=10,
        event_types=[SEARCH_EXTRACT_REQUESTED],
        extractors=[DocxTextExtractor(max_entries=1)],
    )

    async with session_factory() as session:
        version = await session.get(FileVersion, version_id)
    assert version is not None
    assert version.search_status == "skipped"
    assert version.search_text is None
    assert version.search_error == "docx_archive_too_large"
    assert index_adapter.documents == {}


@pytest.mark.asyncio
async def test_search_extract_requested_skips_unsupported_file(
    client: AsyncClient,
    session_factory: async_sessionmaker[AsyncSession],
    settings: Settings,
    storage_adapter: InMemoryStorageAdapter,
) -> None:
    await seed_admin(session_factory, settings)
    token = await login(client)
    space = await create_space(client, token, slug="search-extract-skip-space")
    tenant_id = UUID(str(space["tenant_id"]))
    async with session_factory() as session:
        session.add(
            FileBlob(
                tenant_id=tenant_id,
                hash_algo="sha256",
                content_hash="2" * 64,
                size_bytes=256,
                storage_key="objects/test/image",
                mime_type="image/png",
                ref_count=0,
            )
        )
        await session.commit()

    response = await client.post(
        "/api/v1/uploads/init",
        headers={"X-CSRF-Token": token},
        json={
            "space_id": space["id"],
            "parent_id": space["root_node_id"],
            "file_name": "图片.png",
            "size_bytes": 256,
            "content_hash": "2" * 64,
            "hash_algo": "sha256",
            "mime_type": "image/png",
        },
    )
    assert response.status_code == 201
    version_id = UUID(response.json()["version_id"])
    index_adapter = InMemorySearchIndexAdapter()

    await _dispatch_search_events(
        session_factory=session_factory,
        settings=settings,
        index_adapter=index_adapter,
        storage_adapter=storage_adapter,
        batch_size=10,
        event_types=[SEARCH_EXTRACT_REQUESTED],
    )

    async with session_factory() as session:
        version = await session.get(FileVersion, version_id)
    assert version is not None
    assert version.search_status == "skipped"
    assert version.search_text is None
    assert version.search_error == "unsupported_mime_type"


@pytest.mark.asyncio
async def test_search_extract_requested_skips_large_text_file(
    client: AsyncClient,
    session_factory: async_sessionmaker[AsyncSession],
    settings: Settings,
    storage_adapter: InMemoryStorageAdapter,
) -> None:
    await seed_admin(session_factory, settings)
    token = await login(client)
    space = await create_space(client, token, slug="search-extract-large-text-space")
    tenant_id = UUID(str(space["tenant_id"]))
    storage_key = "objects/test/large-text"
    storage_adapter.object_contents[(settings.s3_bucket, storage_key)] = b"x" * 9
    async with session_factory() as session:
        session.add(
            FileBlob(
                tenant_id=tenant_id,
                hash_algo="sha256",
                content_hash="5" * 64,
                size_bytes=9,
                storage_key=storage_key,
                mime_type="text/plain",
                ref_count=0,
            )
        )
        await session.commit()

    response = await client.post(
        "/api/v1/uploads/init",
        headers={"X-CSRF-Token": token},
        json={
            "space_id": space["id"],
            "parent_id": space["root_node_id"],
            "file_name": "过大文本.txt",
            "size_bytes": 9,
            "content_hash": "5" * 64,
            "hash_algo": "sha256",
            "mime_type": "text/plain",
        },
    )
    assert response.status_code == 201
    version_id = UUID(response.json()["version_id"])
    index_adapter = InMemorySearchIndexAdapter()

    await _dispatch_search_events(
        session_factory=session_factory,
        settings=settings,
        index_adapter=index_adapter,
        storage_adapter=storage_adapter,
        batch_size=10,
        event_types=[SEARCH_EXTRACT_REQUESTED],
        max_bytes=8,
    )

    async with session_factory() as session:
        version = await session.get(FileVersion, version_id)
    assert version is not None
    assert version.search_status == "skipped"
    assert version.search_text is None
    assert version.search_error == "file_too_large"
    assert index_adapter.documents == {}


@pytest.mark.asyncio
async def test_search_extract_requested_skips_oversized_extracted_text(
    client: AsyncClient,
    session_factory: async_sessionmaker[AsyncSession],
    settings: Settings,
    storage_adapter: InMemoryStorageAdapter,
) -> None:
    await seed_admin(session_factory, settings)
    token = await login(client)
    space = await create_space(client, token, slug="search-extract-large-output-space")
    tenant_id = UUID(str(space["tenant_id"]))
    storage_key = "objects/test/small-source-large-output"
    storage_adapter.object_contents[(settings.s3_bucket, storage_key)] = b"source"
    async with session_factory() as session:
        session.add(
            FileBlob(
                tenant_id=tenant_id,
                hash_algo="sha256",
                content_hash="7" * 64,
                size_bytes=6,
                storage_key=storage_key,
                mime_type="text/plain",
                ref_count=0,
            )
        )
        await session.commit()

    response = await client.post(
        "/api/v1/uploads/init",
        headers={"X-CSRF-Token": token},
        json={
            "space_id": space["id"],
            "parent_id": space["root_node_id"],
            "file_name": "抽取结果过大.txt",
            "size_bytes": 6,
            "content_hash": "7" * 64,
            "hash_algo": "sha256",
            "mime_type": "text/plain",
        },
    )
    assert response.status_code == 201
    version_id = UUID(response.json()["version_id"])
    index_adapter = InMemorySearchIndexAdapter()

    await _dispatch_search_events(
        session_factory=session_factory,
        settings=settings,
        index_adapter=index_adapter,
        storage_adapter=storage_adapter,
        batch_size=10,
        event_types=[SEARCH_EXTRACT_REQUESTED],
        max_bytes=8,
        extractors=[OversizedTextExtractor()],
    )

    async with session_factory() as session:
        version = await session.get(FileVersion, version_id)
    assert version is not None
    assert version.search_status == "skipped"
    assert version.search_text is None
    assert version.search_error == "extracted_text_too_large"
    assert index_adapter.documents == {}


@pytest.mark.asyncio
async def test_search_extract_requested_marks_decode_failure_without_retry(
    client: AsyncClient,
    session_factory: async_sessionmaker[AsyncSession],
    settings: Settings,
    storage_adapter: InMemoryStorageAdapter,
) -> None:
    await seed_admin(session_factory, settings)
    token = await login(client)
    space = await create_space(client, token, slug="search-extract-decode-failed-space")
    tenant_id = UUID(str(space["tenant_id"]))
    storage_key = "objects/test/binary-text"
    storage_adapter.object_contents[(settings.s3_bucket, storage_key)] = b"\xff\xfe\x00"
    async with session_factory() as session:
        session.add(
            FileBlob(
                tenant_id=tenant_id,
                hash_algo="sha256",
                content_hash="3" * 64,
                size_bytes=3,
                storage_key=storage_key,
                mime_type="text/plain",
                ref_count=0,
            )
        )
        await session.commit()

    response = await client.post(
        "/api/v1/uploads/init",
        headers={"X-CSRF-Token": token},
        json={
            "space_id": space["id"],
            "parent_id": space["root_node_id"],
            "file_name": "坏编码.txt",
            "size_bytes": 3,
            "content_hash": "3" * 64,
            "hash_algo": "sha256",
            "mime_type": "text/plain",
        },
    )
    assert response.status_code == 201
    version_id = UUID(response.json()["version_id"])
    index_adapter = InMemorySearchIndexAdapter()

    result = await _dispatch_search_events(
        session_factory=session_factory,
        settings=settings,
        index_adapter=index_adapter,
        storage_adapter=storage_adapter,
        batch_size=10,
        event_types=[SEARCH_EXTRACT_REQUESTED],
    )

    assert result.to_dict() == {"claimed": 1, "sent": 1, "failed": 0, "dead": 0}
    async with session_factory() as session:
        version = await session.get(FileVersion, version_id)
        event = (
            await session.execute(
                select(OutboxEvent).where(OutboxEvent.event_type == SEARCH_EXTRACT_REQUESTED)
            )
        ).scalar_one()

    assert version is not None
    assert version.search_status == "failed"
    assert version.search_text is None
    assert version.search_error == "decode_failed"
    assert event.status == "sent"
    assert event.retry_count == 0
    assert index_adapter.documents == {}


@pytest.mark.asyncio
async def test_search_extract_requested_retries_storage_read_failure(
    client: AsyncClient,
    session_factory: async_sessionmaker[AsyncSession],
    settings: Settings,
) -> None:
    await seed_admin(session_factory, settings)
    token = await login(client)
    space = await create_space(client, token, slug="search-extract-storage-failed-space")
    tenant_id = UUID(str(space["tenant_id"]))
    async with session_factory() as session:
        session.add(
            FileBlob(
                tenant_id=tenant_id,
                hash_algo="sha256",
                content_hash="4" * 64,
                size_bytes=8,
                storage_key="objects/test/storage-failed",
                mime_type="text/plain",
                ref_count=0,
            )
        )
        await session.commit()

    response = await client.post(
        "/api/v1/uploads/init",
        headers={"X-CSRF-Token": token},
        json={
            "space_id": space["id"],
            "parent_id": space["root_node_id"],
            "file_name": "存储失败.txt",
            "size_bytes": 8,
            "content_hash": "4" * 64,
            "hash_algo": "sha256",
            "mime_type": "text/plain",
        },
    )
    assert response.status_code == 201
    version_id = UUID(response.json()["version_id"])
    index_adapter = InMemorySearchIndexAdapter()

    result = await _dispatch_search_events(
        session_factory=session_factory,
        settings=settings,
        index_adapter=index_adapter,
        storage_adapter=FailingReadStorageAdapter(),
        batch_size=10,
        event_types=[SEARCH_EXTRACT_REQUESTED],
        expected_failed=1,
    )

    assert result.to_dict() == {"claimed": 1, "sent": 0, "failed": 1, "dead": 0}
    async with session_factory() as session:
        version = await session.get(FileVersion, version_id)
        event = (
            await session.execute(
                select(OutboxEvent).where(OutboxEvent.event_type == SEARCH_EXTRACT_REQUESTED)
            )
        ).scalar_one()

    assert version is not None
    assert version.search_status == "failed"
    assert version.search_text is None
    assert version.search_error == "storage_read_failed"
    assert event.status == "failed"
    assert event.retry_count == 1
    assert index_adapter.documents == {}


@pytest.mark.asyncio
async def test_file_change_search_events_update_and_delete_index_document(
    client: AsyncClient,
    session_factory: async_sessionmaker[AsyncSession],
    settings: Settings,
) -> None:
    await seed_admin(session_factory, settings)
    token = await login(client)
    space = await create_space(client, token, slug="search-file-change-events")
    tenant_id = UUID(str(space["tenant_id"]))
    async with session_factory() as session:
        session.add(
            FileBlob(
                tenant_id=tenant_id,
                hash_algo="sha256",
                content_hash="f" * 64,
                size_bytes=384,
                storage_key="objects/test/ff",
                mime_type="text/plain",
                ref_count=0,
            )
        )
        await session.commit()

    upload_response = await client.post(
        "/api/v1/uploads/init",
        headers={"X-CSRF-Token": token},
        json={
            "space_id": space["id"],
            "parent_id": space["root_node_id"],
            "file_name": "索引变更前.txt",
            "size_bytes": 384,
            "content_hash": "f" * 64,
            "hash_algo": "sha256",
            "mime_type": "text/plain",
        },
    )
    assert upload_response.status_code == 201
    node_id = UUID(upload_response.json()["node_id"])
    index_adapter = InMemorySearchIndexAdapter()

    await _dispatch_search_events(
        session_factory=session_factory,
        settings=settings,
        index_adapter=index_adapter,
        batch_size=10,
    )
    assert index_adapter.documents[f"{tenant_id}:{node_id}"].name == "索引变更前.txt"

    rename_response = await client.patch(
        f"/api/v1/files/{node_id}",
        headers={"X-CSRF-Token": token},
        json={"name": "索引变更后.txt"},
    )
    assert rename_response.status_code == 200

    await _dispatch_search_events(
        session_factory=session_factory,
        settings=settings,
        index_adapter=index_adapter,
        batch_size=10,
    )
    assert index_adapter.documents[f"{tenant_id}:{node_id}"].name == "索引变更后.txt"

    delete_response = await client.delete(
        f"/api/v1/files/{node_id}",
        headers={"X-CSRF-Token": token},
    )
    assert delete_response.status_code == 200

    await _dispatch_search_events(
        session_factory=session_factory,
        settings=settings,
        index_adapter=index_adapter,
        batch_size=10,
    )
    assert f"{tenant_id}:{node_id}" not in index_adapter.documents
    assert index_adapter.deleted_document_ids[-1] == f"{tenant_id}:{node_id}"


@pytest.mark.asyncio
async def test_acl_rebuild_event_reindexes_file_acl_tokens(
    client: AsyncClient,
    session_factory: async_sessionmaker[AsyncSession],
    settings: Settings,
) -> None:
    await seed_admin(session_factory, settings)
    token = await login(client)
    space = await create_space(client, token, slug="search-acl-rebuild-space")
    tenant_id = UUID(str(space["tenant_id"]))
    _, actor_id = await _get_seeded_tenant_and_user(session_factory)
    subject_id = uuid4()
    async with session_factory() as session:
        session.add(
            FileBlob(
                tenant_id=tenant_id,
                hash_algo="sha256",
                content_hash="e" * 64,
                size_bytes=512,
                storage_key="objects/test/ee",
                mime_type="text/plain",
                ref_count=0,
            )
        )
        await session.commit()

    response = await client.post(
        "/api/v1/uploads/init",
        headers={"X-CSRF-Token": token},
        json={
            "space_id": space["id"],
            "parent_id": space["root_node_id"],
            "file_name": "ACL重建.txt",
            "size_bytes": 512,
            "content_hash": "e" * 64,
            "hash_algo": "sha256",
            "mime_type": "text/plain",
        },
    )
    assert response.status_code == 201
    node_id = UUID(response.json()["node_id"])
    index_adapter = InMemorySearchIndexAdapter()

    async with session_factory() as session:
        dispatcher = OutboxDispatcher(
            repository=AuditRepository(session),
            publisher=search_tasks.SearchOutboxPublisher(
                index_service=SearchIndexService(
                    repository=SearchRepository(session),
                    index_adapter=index_adapter,
                ),
                fallback_publisher=LoggingOutboxPublisher(),
            ),
            max_retries=settings.outbox_max_retries,
        )
        await dispatcher.dispatch_pending(batch_size=10, event_types=[SEARCH_INDEX_REQUESTED])
        await session.commit()

    document_id = f"{tenant_id}:{node_id}"
    assert index_adapter.documents[document_id].acl_tokens == [f"space:{space['id']}:role:owner"]

    async with session_factory() as session:
        session.add(
            AclEntry(
                tenant_id=tenant_id,
                node_id=node_id,
                subject_type="user",
                subject_id=subject_id,
                effect="allow",
                actions=["download"],
                inherit=True,
                created_by=actor_id,
            )
        )
        await AuditRepository(session).add_outbox_event(
            tenant_id=tenant_id,
            event_type=SEARCH_ACL_REBUILD_REQUESTED,
            aggregate_type="node",
            aggregate_id=node_id,
            payload={
                "scope": "node",
                "resource_id": str(node_id),
                "permission_version": 2,
                "reason": "node_acl_created",
            },
        )
        await session.commit()

    async with session_factory() as session:
        dispatcher = OutboxDispatcher(
            repository=AuditRepository(session),
            publisher=search_tasks.SearchOutboxPublisher(
                index_service=SearchIndexService(
                    repository=SearchRepository(session),
                    index_adapter=index_adapter,
                ),
                fallback_publisher=LoggingOutboxPublisher(),
            ),
            max_retries=settings.outbox_max_retries,
        )
        result = await dispatcher.dispatch_pending(
            batch_size=10,
            event_types=[SEARCH_ACL_REBUILD_REQUESTED],
        )
        await session.commit()

    assert result.to_dict() == {"claimed": 1, "sent": 1, "failed": 0, "dead": 0}
    rebuilt_document = index_adapter.documents[document_id]
    assert rebuilt_document.acl_tokens == sorted(
        [f"space:{space['id']}:role:owner", f"user:{subject_id}"]
    )


async def _dispatch_search_events(
    *,
    session_factory: async_sessionmaker[AsyncSession],
    settings: Settings,
    index_adapter: InMemorySearchIndexAdapter,
    batch_size: int,
    storage_adapter: InMemoryStorageAdapter | None = None,
    event_types: list[str] | None = None,
    expected_failed: int = 0,
    expected_dead: int = 0,
    max_bytes: int | None = None,
    extractors: list[TextExtractor] | None = None,
) -> OutboxDispatchResult:
    async with session_factory() as session:
        repository = SearchRepository(session)
        index_service = SearchIndexService(
            repository=repository,
            index_adapter=index_adapter,
        )
        dispatcher = OutboxDispatcher(
            repository=AuditRepository(session),
            publisher=search_tasks.SearchOutboxPublisher(
                index_service=index_service,
                extraction_service=(
                    SearchExtractionService(
                        repository=repository,
                        index_service=index_service,
                        storage=storage_adapter,
                        settings=settings,
                        max_bytes=max_bytes,
                        extractors=extractors,
                    )
                    if storage_adapter is not None
                    else None
                ),
                fallback_publisher=LoggingOutboxPublisher(),
            ),
            max_retries=settings.outbox_max_retries,
        )
        result = await dispatcher.dispatch_pending(
            batch_size=batch_size,
            event_types=event_types or [SEARCH_INDEX_REQUESTED],
        )
        await session.commit()
    assert result.failed == expected_failed
    assert result.dead == expected_dead
    return result


async def _upload_existing_blob_and_extract_text(
    *,
    client: AsyncClient,
    session_factory: async_sessionmaker[AsyncSession],
    settings: Settings,
    storage_adapter: InMemoryStorageAdapter,
    slug: str,
    storage_key: str,
    file_name: str,
    content_hash: str,
    mime_type: str,
    content: bytes,
) -> tuple[UUID, str, UUID, InMemorySearchIndexAdapter]:
    await seed_admin(session_factory, settings)
    token = await login(client)
    space = await create_space(client, token, slug=slug)
    tenant_id = UUID(str(space["tenant_id"]))
    storage_adapter.object_contents[(settings.s3_bucket, storage_key)] = content
    async with session_factory() as session:
        session.add(
            FileBlob(
                tenant_id=tenant_id,
                hash_algo="sha256",
                content_hash=content_hash,
                size_bytes=len(content),
                storage_key=storage_key,
                mime_type=mime_type,
                ref_count=0,
            )
        )
        await session.commit()

    response = await client.post(
        "/api/v1/uploads/init",
        headers={"X-CSRF-Token": token},
        json={
            "space_id": space["id"],
            "parent_id": space["root_node_id"],
            "file_name": file_name,
            "size_bytes": len(content),
            "content_hash": content_hash,
            "hash_algo": "sha256",
            "mime_type": mime_type,
        },
    )
    assert response.status_code == 201
    node_id = response.json()["node_id"]
    version_id = UUID(response.json()["version_id"])
    index_adapter = InMemorySearchIndexAdapter()

    await _dispatch_search_events(
        session_factory=session_factory,
        settings=settings,
        index_adapter=index_adapter,
        storage_adapter=storage_adapter,
        batch_size=10,
        event_types=[SEARCH_EXTRACT_REQUESTED],
    )
    return tenant_id, node_id, version_id, index_adapter


def _build_pdf_with_text(*texts: str) -> bytes:
    writer = PdfWriter()
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    font_ref = writer._add_object(font)
    for text in texts:
        page = writer.add_blank_page(width=300, height=300)
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font_ref})}
        )
        stream = DecodedStreamObject()
        stream.set_data(f"BT /F1 12 Tf 50 250 Td ({text}) Tj ET".encode())
        page[NameObject("/Contents")] = writer._add_object(stream)
    buffer = BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def _build_docx_with_text(paragraph_text: str, table_row: tuple[str, str]) -> bytes:
    document = Document()
    document.add_paragraph(paragraph_text)
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = table_row[0]
    table.cell(0, 1).text = table_row[1]
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _build_pptx_with_text(shape_text: str, table_row: tuple[str, str]) -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    text_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(4), Inches(1))
    text_box.text = shape_text
    table = slide.shapes.add_table(
        rows=1,
        cols=2,
        left=Inches(0.5),
        top=Inches(1.5),
        width=Inches(4),
        height=Inches(1),
    ).table
    table.cell(0, 0).text = table_row[0]
    table.cell(0, 1).text = table_row[1]
    buffer = BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def _build_xlsx_with_text(first_cell: str, table_row: tuple[str, str]) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet["A1"] = first_cell
    worksheet["A2"] = table_row[0]
    worksheet["B2"] = table_row[1]
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()
