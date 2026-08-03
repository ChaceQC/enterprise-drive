from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from typing import Any, Protocol
from zipfile import BadZipFile, ZipFile

from docx import Document
from docx.opc.exceptions import PackageNotFoundError as DocxPackageNotFoundError
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pptx import Presentation
from pptx.exc import PackageNotFoundError as PptxPackageNotFoundError
from pypdf import PdfReader
from pypdf.errors import PdfReadError

_TEXT_MIME_TYPES = {
    "application/csv",
    "application/json",
    "application/ld+json",
    "application/xml",
    "application/yaml",
    "text/csv",
    "text/markdown",
    "text/plain",
    "text/tab-separated-values",
    "text/xml",
}
_TEXT_EXTENSIONS = {
    ".csv",
    ".json",
    ".log",
    ".md",
    ".markdown",
    ".txt",
    ".tsv",
    ".xml",
    ".yaml",
    ".yml",
}
_PDF_MIME_TYPES = {"application/pdf"}
_PDF_EXTENSIONS = {".pdf"}
_DEFAULT_PDF_MAX_PAGES = 50
_IMAGE_MIME_TYPES = {
    "image/bmp",
    "image/gif",
    "image/jpeg",
    "image/png",
    "image/tiff",
    "image/webp",
}
_IMAGE_EXTENSIONS = {
    ".bmp",
    ".gif",
    ".jpeg",
    ".jpg",
    ".png",
    ".tif",
    ".tiff",
    ".webp",
}
_LEGACY_OFFICE_MIME_TYPES = {
    "application/msword",
    "application/rtf",
    "application/vnd.ms-excel",
    "application/vnd.ms-powerpoint",
    "application/vnd.oasis.opendocument.presentation",
    "application/vnd.oasis.opendocument.spreadsheet",
    "application/vnd.oasis.opendocument.text",
    "text/rtf",
}
_LEGACY_OFFICE_EXTENSIONS = {
    ".doc",
    ".odp",
    ".ods",
    ".odt",
    ".ppt",
    ".rtf",
    ".xls",
}
_DOCX_MIME_TYPES = {
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}
_DOCX_EXTENSIONS = {".docx"}
_DEFAULT_DOCX_MAX_UNCOMPRESSED_BYTES = 5 * 1024 * 1024
_DEFAULT_DOCX_MAX_ENTRIES = 256
_PPTX_MIME_TYPES = {
    "application/vnd.openxmlformats-officedocument.presentationml.presentation",
}
_PPTX_EXTENSIONS = {".pptx"}
_DEFAULT_PPTX_MAX_UNCOMPRESSED_BYTES = 5 * 1024 * 1024
_DEFAULT_PPTX_MAX_ENTRIES = 256
_XLSX_MIME_TYPES = {
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}
_XLSX_EXTENSIONS = {".xlsx"}
_DEFAULT_XLSX_MAX_UNCOMPRESSED_BYTES = 5 * 1024 * 1024
_DEFAULT_XLSX_MAX_ENTRIES = 512


class TextExtractionError(Exception):
    def __init__(self, reason: str, *, status: str = "failed") -> None:
        super().__init__(reason)
        self.reason = reason
        self.status = status


@dataclass(frozen=True)
class TextExtractionContext:
    mime_type: str
    name: str


class TextExtractor(Protocol):
    def supports(self, context: TextExtractionContext) -> bool: ...

    def extract(self, content: bytes, context: TextExtractionContext) -> str: ...


class OcrTextEngine(Protocol):
    def extract_image(self, content: bytes, context: TextExtractionContext) -> str: ...

    def extract_pdf(
        self,
        content: bytes,
        context: TextExtractionContext,
        *,
        max_pages: int,
    ) -> str: ...


class OfficeDocumentConverter(Protocol):
    def convert_to_pdf(self, content: bytes, context: TextExtractionContext) -> bytes: ...


class Utf8TextExtractor:
    def supports(self, context: TextExtractionContext) -> bool:
        if context.mime_type.startswith("text/") or context.mime_type in _TEXT_MIME_TYPES:
            return True
        normalized_name = context.name.casefold()
        return any(normalized_name.endswith(extension) for extension in _TEXT_EXTENSIONS)

    def extract(self, content: bytes, context: TextExtractionContext) -> str:
        try:
            return content.decode("utf-8-sig").replace("\x00", "")
        except UnicodeDecodeError as exc:
            raise TextExtractionError("decode_failed") from exc


class PdfTextExtractor:
    def __init__(
        self,
        *,
        max_pages: int = _DEFAULT_PDF_MAX_PAGES,
        ocr_engine: OcrTextEngine | None = None,
        max_source_bytes: int | None = None,
    ) -> None:
        self.max_pages = max_pages
        self.ocr_engine = ocr_engine
        self.max_source_bytes = max_source_bytes

    def supports(self, context: TextExtractionContext) -> bool:
        normalized_name = context.name.casefold()
        return context.mime_type in _PDF_MIME_TYPES or any(
            normalized_name.endswith(extension) for extension in _PDF_EXTENSIONS
        )

    def extract(self, content: bytes, context: TextExtractionContext) -> str:
        try:
            reader = PdfReader(BytesIO(content), strict=False)
            if reader.is_encrypted:
                raise TextExtractionError("pdf_encrypted")
            page_texts: list[str] = []
            for page_number, page in enumerate(reader.pages):
                if page_number >= self.max_pages:
                    break
                page_texts.append(page.extract_text() or "")
        except TextExtractionError:
            raise
        except (PdfReadError, UnicodeDecodeError, OSError, ValueError) as exc:
            raise TextExtractionError("decode_failed") from exc
        text = "\n".join(text.strip() for text in page_texts if text.strip())
        if text or self.ocr_engine is None:
            return text
        return self.ocr_engine.extract_pdf(
            content,
            context,
            max_pages=self.max_pages,
        )


class ImageOcrTextExtractor:
    def __init__(
        self,
        *,
        ocr_engine: OcrTextEngine,
        max_source_bytes: int | None = None,
    ) -> None:
        self.ocr_engine = ocr_engine
        self.max_source_bytes = max_source_bytes

    def supports(self, context: TextExtractionContext) -> bool:
        normalized_name = context.name.casefold()
        return context.mime_type in _IMAGE_MIME_TYPES or any(
            normalized_name.endswith(extension) for extension in _IMAGE_EXTENSIONS
        )

    def extract(self, content: bytes, context: TextExtractionContext) -> str:
        return self.ocr_engine.extract_image(content, context)


class OfficeDocumentTextExtractor:
    def __init__(
        self,
        *,
        converter: OfficeDocumentConverter,
        pdf_extractor: PdfTextExtractor,
        max_source_bytes: int | None = None,
    ) -> None:
        self.converter = converter
        self.pdf_extractor = pdf_extractor
        self.max_source_bytes = max_source_bytes

    def supports(self, context: TextExtractionContext) -> bool:
        normalized_name = context.name.casefold()
        return context.mime_type in _LEGACY_OFFICE_MIME_TYPES or any(
            normalized_name.endswith(extension) for extension in _LEGACY_OFFICE_EXTENSIONS
        )

    def extract(self, content: bytes, context: TextExtractionContext) -> str:
        pdf_content = self.converter.convert_to_pdf(content, context)
        return self.pdf_extractor.extract(
            pdf_content,
            TextExtractionContext(
                mime_type="application/pdf",
                name=f"{context.name}.pdf",
            ),
        )


class DocxTextExtractor:
    def __init__(
        self,
        *,
        max_uncompressed_bytes: int = _DEFAULT_DOCX_MAX_UNCOMPRESSED_BYTES,
        max_entries: int = _DEFAULT_DOCX_MAX_ENTRIES,
    ) -> None:
        self.max_uncompressed_bytes = max_uncompressed_bytes
        self.max_entries = max_entries

    def supports(self, context: TextExtractionContext) -> bool:
        normalized_name = context.name.casefold()
        return context.mime_type in _DOCX_MIME_TYPES or any(
            normalized_name.endswith(extension) for extension in _DOCX_EXTENSIONS
        )

    def extract(self, content: bytes, context: TextExtractionContext) -> str:
        try:
            _check_archive_limits(
                content,
                max_uncompressed_bytes=self.max_uncompressed_bytes,
                max_entries=self.max_entries,
                reason="docx_archive_too_large",
            )
            document = Document(BytesIO(content))
        except (BadZipFile, KeyError, DocxPackageNotFoundError, ValueError) as exc:
            raise TextExtractionError("decode_failed") from exc
        paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs]
        table_cells = [
            cell.text.strip()
            for table in document.tables
            for row in table.rows
            for cell in row.cells
        ]
        return "\n".join(text for text in [*paragraphs, *table_cells] if text)


class PptxTextExtractor:
    def __init__(
        self,
        *,
        max_uncompressed_bytes: int = _DEFAULT_PPTX_MAX_UNCOMPRESSED_BYTES,
        max_entries: int = _DEFAULT_PPTX_MAX_ENTRIES,
    ) -> None:
        self.max_uncompressed_bytes = max_uncompressed_bytes
        self.max_entries = max_entries

    def supports(self, context: TextExtractionContext) -> bool:
        normalized_name = context.name.casefold()
        return context.mime_type in _PPTX_MIME_TYPES or any(
            normalized_name.endswith(extension) for extension in _PPTX_EXTENSIONS
        )

    def extract(self, content: bytes, context: TextExtractionContext) -> str:
        try:
            _check_archive_limits(
                content,
                max_uncompressed_bytes=self.max_uncompressed_bytes,
                max_entries=self.max_entries,
                reason="pptx_archive_too_large",
            )
            presentation = Presentation(BytesIO(content))
        except (BadZipFile, KeyError, PptxPackageNotFoundError, ValueError) as exc:
            raise TextExtractionError("decode_failed") from exc
        texts = [
            text
            for slide in presentation.slides
            for text in self._extract_shape_texts(slide.shapes)
        ]
        return "\n".join(text for text in texts if text)

    def _extract_shape_texts(self, shapes: Any) -> list[str]:
        texts: list[str] = []
        for shape in shapes:
            if getattr(shape, "has_table", False):
                table = shape.table
                texts.extend(
                    cell.text.strip()
                    for row in table.rows
                    for cell in row.cells
                    if cell.text.strip()
                )
                continue
            shape_text = getattr(shape, "text", "")
            if isinstance(shape_text, str) and shape_text.strip():
                texts.append(shape_text.strip())
                continue
            nested_shapes = getattr(shape, "shapes", None)
            if nested_shapes is not None:
                texts.extend(self._extract_shape_texts(nested_shapes))
        return texts


class XlsxTextExtractor:
    def __init__(
        self,
        *,
        max_uncompressed_bytes: int = _DEFAULT_XLSX_MAX_UNCOMPRESSED_BYTES,
        max_entries: int = _DEFAULT_XLSX_MAX_ENTRIES,
    ) -> None:
        self.max_uncompressed_bytes = max_uncompressed_bytes
        self.max_entries = max_entries

    def supports(self, context: TextExtractionContext) -> bool:
        normalized_name = context.name.casefold()
        return context.mime_type in _XLSX_MIME_TYPES or any(
            normalized_name.endswith(extension) for extension in _XLSX_EXTENSIONS
        )

    def extract(self, content: bytes, context: TextExtractionContext) -> str:
        workbook: Any | None = None
        try:
            _check_archive_limits(
                content,
                max_uncompressed_bytes=self.max_uncompressed_bytes,
                max_entries=self.max_entries,
                reason="xlsx_archive_too_large",
            )
            workbook = load_workbook(
                filename=BytesIO(content),
                read_only=True,
                data_only=True,
            )
            texts = [
                str(value).strip()
                for worksheet in workbook.worksheets
                for row in worksheet.iter_rows(values_only=True)
                for value in row
                if value is not None and str(value).strip()
            ]
        except (BadZipFile, InvalidFileException, KeyError, OSError, ValueError) as exc:
            raise TextExtractionError("decode_failed") from exc
        finally:
            if workbook is not None:
                workbook.close()
        return "\n".join(texts)


def _check_archive_limits(
    content: bytes,
    *,
    max_uncompressed_bytes: int,
    max_entries: int,
    reason: str,
) -> None:
    with ZipFile(BytesIO(content)) as archive:
        entries = archive.infolist()
        if len(entries) > max_entries:
            raise TextExtractionError(reason, status="skipped")
        uncompressed_size = sum(entry.file_size for entry in entries)
        if uncompressed_size > max_uncompressed_bytes:
            raise TextExtractionError(reason, status="skipped")


def default_text_extractors(
    *,
    ocr_engine: OcrTextEngine | None = None,
    office_converter: OfficeDocumentConverter | None = None,
    complex_source_max_bytes: int | None = None,
    pdf_max_pages: int = _DEFAULT_PDF_MAX_PAGES,
) -> list[TextExtractor]:
    pdf_extractor = PdfTextExtractor(
        max_pages=pdf_max_pages,
        ocr_engine=ocr_engine,
        max_source_bytes=complex_source_max_bytes if ocr_engine is not None else None,
    )
    extractors: list[TextExtractor] = [
        Utf8TextExtractor(),
        pdf_extractor,
        DocxTextExtractor(),
        PptxTextExtractor(),
        XlsxTextExtractor(),
    ]
    if ocr_engine is not None:
        extractors.append(
            ImageOcrTextExtractor(
                ocr_engine=ocr_engine,
                max_source_bytes=complex_source_max_bytes,
            )
        )
    if office_converter is not None:
        extractors.append(
            OfficeDocumentTextExtractor(
                converter=office_converter,
                pdf_extractor=pdf_extractor,
                max_source_bytes=complex_source_max_bytes,
            )
        )
    return extractors
